# -*- coding: utf-8 -*-
"""Trích văn bản an toàn từ tài liệu tạm của một lượt AI Search."""
import base64
import csv
import io
import json
import zipfile
from pathlib import Path
from xml.etree import ElementTree

MAX_FILES = 5
MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ARCHIVE_BYTES = 50 * 1024 * 1024
MAX_CHARS_PER_FILE = 30_000
MAX_TOTAL_CHARS = 60_000
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp"}
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".xlsx", ".pptx", ".txt", ".md", ".csv", ".json"} | IMAGE_EXTENSIONS


class AttachmentError(ValueError):
    pass


def extract_uploads(files):
    uploads = list(files)
    if len(uploads) > MAX_FILES:
        raise AttachmentError(f"Chỉ được đính kèm tối đa {MAX_FILES} tệp.")
    sections = []
    metadata = []
    total = 0
    for upload in uploads:
        name = Path(str(upload.name or "")).name
        extension = Path(name).suffix.lower()
        if extension not in ALLOWED_EXTENSIONS:
            raise AttachmentError(f"Không hỗ trợ định dạng {extension or '(không có đuôi)'}.")
        if upload.size > MAX_FILE_BYTES:
            raise AttachmentError(f"Tệp {name} vượt quá giới hạn 10 MB.")
        try:
            text = extract_file(upload).strip()
        except AttachmentError:
            raise
        except Exception as exc:  # noqa: BLE001 - parser của từng định dạng có nhiều loại lỗi
            raise AttachmentError(f"Không đọc được tệp {name}: tệp có thể hỏng hoặc được bảo vệ.") from exc
        if not text:
            raise AttachmentError(f"Không đọc được nội dung chữ trong tệp {name}.")
        remaining = MAX_TOTAL_CHARS - total
        if remaining <= 0:
            break
        text = text[:min(MAX_CHARS_PER_FILE, remaining)]
        total += len(text)
        sections.append(f"--- Tài liệu: {name} ---\n{text}")
        metadata.append({"name": name, "size": upload.size, "characters": len(text)})
    return "\n\n".join(sections), metadata


def _extract(upload, extension):
    upload.seek(0)
    if extension in IMAGE_EXTENSIONS:
        return _extract_image(upload)
    if extension in {".txt", ".md", ".csv", ".json"}:
        raw = upload.read(MAX_FILE_BYTES + 1)
        text = raw.decode("utf-8-sig", errors="replace")
        if extension == ".json":
            try:
                return json.dumps(json.loads(text), ensure_ascii=False)
            except json.JSONDecodeError:
                return text
        if extension == ".csv":
            return "\n".join(" | ".join(row) for row in csv.reader(io.StringIO(text)))
        return text
    if extension == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(upload)
        return "\n".join((page.extract_text() or "") for page in reader.pages[:50])
    if extension == ".docx":
        _validate_archive(upload)
        from docx import Document
        document = Document(upload)
        paragraphs = [p.text for p in document.paragraphs]
        for table in document.tables:
            paragraphs.extend(" | ".join(cell.text for cell in row.cells) for row in table.rows)
        return "\n".join(paragraphs)
    if extension == ".xlsx":
        _validate_archive(upload)
        from openpyxl import load_workbook
        workbook = load_workbook(upload, read_only=True, data_only=True)
        lines = []
        for sheet in workbook.worksheets[:10]:
            lines.append(f"[{sheet.title}]")
            for row in sheet.iter_rows(max_row=2000, max_col=50, values_only=True):
                lines.append(" | ".join("" if value is None else str(value) for value in row))
                if sum(map(len, lines)) >= MAX_CHARS_PER_FILE:
                    break
        workbook.close()
        return "\n".join(lines)
    return _pptx_text(upload)


def _extract_image(upload):
    """Đọc ảnh bằng đúng task vision OCR đã được route và audit trên Hub."""
    data = upload.read(MAX_FILE_BYTES + 1)
    if len(data) > MAX_FILE_BYTES:
        raise AttachmentError(f"Tệp {upload.name} vượt quá giới hạn 10 MB.")
    try:
        from PIL import Image
        image = Image.open(io.BytesIO(data)).convert("RGB")
        image.thumbnail((1800, 2400))
        output = io.BytesIO()
        image.save(output, format="JPEG", quality=85, optimize=True)
    except Exception as exc:  # noqa: BLE001
        raise AttachmentError(f"Không đọc được ảnh {upload.name}.") from exc

    from ai.router import complete
    content = [
        {"type": "text", "text": (
            "Đọc chính xác toàn bộ chữ và dữ kiện nhìn thấy trong ảnh. Trả về "
            "văn bản thuần UTF-8, giữ tên riêng, thời gian, chức danh, kỹ năng "
            "và thứ tự mục. Không suy diễn và không thêm dữ kiện.")},
        {"type": "image_url", "image_url": {
            "url": "data:image/jpeg;base64," + base64.b64encode(output.getvalue()).decode("ascii"),
            "detail": "high"}},
    ]
    try:
        completion = complete([
            {"role": "system", "content": "Bạn là OCR chính xác cho tài liệu tuyển dụng."},
            {"role": "user", "content": content},
        ], task="cv_ocr", max_tokens=12_000, reasoning_effort="none")
    except Exception as exc:  # noqa: BLE001 - đổi lỗi provider thành lỗi có ích cho UI
        raise AttachmentError(
            f"Radar chưa đọc được ảnh {upload.name}. Bạn có thể thử lại hoặc gửi PDF/text.") from exc
    text = str(getattr(completion, "text", "") or "").strip()
    if text.startswith("```"):
        text = text.split("\n", 1)[-1]
        if text.endswith("```"):
            text = text[:-3]
    if len(text.strip()) < 10:
        raise AttachmentError(f"Không đọc được đủ nội dung chữ trong ảnh {upload.name}.")
    return text.strip()


def extract_file(upload, max_bytes=MAX_FILE_BYTES):
    """Trích text thuần từ một file; dùng chung cho upload AI Search và CV Hub."""
    extension = Path(str(upload.name or "")).suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise AttachmentError(f"Không hỗ trợ định dạng {extension or '(không có đuôi)'}.")
    if upload.size > max_bytes:
        raise AttachmentError(f"Tệp {upload.name} vượt quá giới hạn đọc {max_bytes // 1024 // 1024} MB.")
    return _extract(upload, extension)


def _pptx_text(upload):
    lines = []
    with zipfile.ZipFile(upload) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_ARCHIVE_BYTES:
            raise AttachmentError("Tệp trình chiếu sau giải nén vượt quá giới hạn an toàn.")
        slides = sorted(name for name in archive.namelist()
                        if name.startswith("ppt/slides/slide") and name.endswith(".xml"))[:100]
        for name in slides:
            root = ElementTree.fromstring(archive.read(name))
            lines.extend(node.text or "" for node in root.iter()
                         if node.tag.endswith("}t") and node.text)
    return "\n".join(lines)


def _validate_archive(upload):
    upload.seek(0)
    with zipfile.ZipFile(upload) as archive:
        if sum(item.file_size for item in archive.infolist()) > MAX_ARCHIVE_BYTES:
            raise AttachmentError("Tệp sau giải nén vượt quá giới hạn an toàn.")
    upload.seek(0)
