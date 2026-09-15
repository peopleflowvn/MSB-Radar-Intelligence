import json
import os
import tempfile
import unittest
import threading
import zipfile

from app.cv_parser import PARSER_VERSION, extract_cv
from app.db import Database, DONE


def _docx_with_textbox(path, visible, textbox_text):
    """Tạo .docx có một đoạn nhìn thấy và một text box - python-docx bỏ sót text box."""
    from docx import Document

    document = Document()
    document.add_paragraph(visible)
    document.save(path)

    injected = (
        "<w:p><w:r><w:txbxContent><w:p><w:r><w:t xml:space=\"preserve\">"
        f"{textbox_text}</w:t></w:r></w:p></w:txbxContent></w:r></w:p>"
    )
    with zipfile.ZipFile(path) as archive:
        items = {name: archive.read(name) for name in archive.namelist()}
    body = items["word/document.xml"].decode("utf-8")
    if "<w:sectPr" in body:
        body = body.replace("<w:sectPr", injected + "<w:sectPr", 1)
    else:
        body = body.replace("</w:body>", injected + "</w:body>", 1)
    items["word/document.xml"] = body.encode("utf-8")
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in items.items():
            archive.writestr(name, data)


class CvParsingTests(unittest.TestCase):
    def test_text_extraction_and_document_full_text_search(self):
        with tempfile.TemporaryDirectory() as tmp:
            cv_path = os.path.join(tmp, "candidate.txt")
            with open(cv_path, "w", encoding="utf-8") as stream:
                stream.write("Kỹ sư dữ liệu Python Apache Airflow")
            result = extract_cv(cv_path)
            self.assertEqual(result["status"], "done")
            self.assertIn("Apache Airflow", result["text"])

            db = Database(os.path.join(tmp, "test.db"), log=lambda *_: None).open()
            try:
                identity = {"source": "topcv", "account": "a@test.vn", "cv_id": "1"}
                db.upsert({**identity, "fullname": "Nguyễn An", "filename": "candidate.txt",
                           "dl_status": DONE})
                db.enqueue_document(**identity, filename="candidate.txt",
                                    parser_version=PARSER_VERSION)
                row = db.claim_document()
                self.assertEqual(row["parse_status"], "pending")
                db.finish_document(row, result)
                found, total = db.query(search='"Apache Airflow"')
                self.assertEqual(total, 1)
                self.assertEqual(found[0]["cv_id"], "1")
            finally:
                db.close()

    def test_backfill_is_idempotent_and_recovers_interrupted_work(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(os.path.join(tmp, "test.db"), log=lambda *_: None).open()
            try:
                db.upsert({"source": "careerviet", "account": "a@test.vn", "cv_id": "99",
                           "fullname": "Backdate", "filename": "99.pdf", "dl_status": DONE})
                db.enqueue_unparsed_documents(PARSER_VERSION)
                row = db.claim_document()
                self.assertIsNotNone(row)
                db.recover_document_queue()
                recovered = db.claim_document()
                self.assertEqual(recovered["cv_id"], "99")
            finally:
                db.close()

    def test_boolean_can_match_metadata_and_cv_text_together(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(os.path.join(tmp, "test.db"), log=lambda *_: None).open()
            try:
                identity = {"source": "topcv", "account": "a@test.vn", "cv_id": "1"}
                db.upsert({**identity, "position": "Java Developer", "filename": "1.txt",
                           "dl_status": DONE})
                db.enqueue_document(**identity, filename="1.txt")
                db.finish_document(db.claim_document(), {
                    "status": "done", "text": "Kubernetes Docker", "method": "text",
                    "quality": 1, "fields": {}, "file_size": 10, "file_mtime_ns": 1})
                rows, total = db.query(search="Java AND Kubernetes")
                self.assertEqual(total, 1)
                self.assertEqual(rows[0]["cv_id"], "1")
            finally:
                db.close()

    def test_delete_cascades_parsing_data(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(os.path.join(tmp, "test.db"), log=lambda *_: None).open()
            try:
                identity = {"source": "topcv", "account": "a", "cv_id": "1"}
                db.upsert({**identity, "filename": "1.txt", "dl_status": DONE})
                db.enqueue_document(**identity, filename="1.txt")
                db.finish_document(db.claim_document(), {"status": "done", "text": "Python",
                    "method": "text", "quality": 1, "fields": {}})
                db.delete_candidates(source="topcv")
                self.assertEqual(db.conn.execute(
                    "SELECT COUNT(*) FROM candidate_documents").fetchone()[0], 0)
                self.assertEqual(db.conn.execute(
                    "SELECT COUNT(*) FROM candidate_extensions").fetchone()[0], 0)
                self.assertEqual(db.conn.execute(
                    "SELECT COUNT(*) FROM candidate_search").fetchone()[0], 0)
            finally:
                db.close()

    def test_claim_is_atomic_across_connections(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "test.db")
            seed = Database(path, log=lambda *_: None).open()
            for cv_id in ("1", "2"):
                seed.upsert({"source": "topcv", "account": "a", "cv_id": cv_id,
                             "filename": cv_id + ".txt", "dl_status": DONE})
                seed.enqueue_document("topcv", "a", cv_id, cv_id + ".txt")
            seed.close()
            claimed, barrier = [], threading.Barrier(2)
            def worker():
                db = Database(path, log=lambda *_: None).open()
                barrier.wait()
                row = db.claim_document()
                claimed.append(row["cv_id"] if row else None)
                db.close()
            threads = [threading.Thread(target=worker) for _ in range(2)]
            for thread in threads: thread.start()
            for thread in threads: thread.join()
            self.assertEqual(set(claimed), {"1", "2"})

    def test_finish_is_atomic_when_extension_payload_fails(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(os.path.join(tmp, "test.db"), log=lambda *_: None).open()
            try:
                identity = {"source": "topcv", "account": "a", "cv_id": "1"}
                db.upsert({**identity, "filename": "1.txt", "dl_status": DONE})
                db.enqueue_document(**identity, filename="1.txt")
                row = db.claim_document()
                with self.assertRaises(TypeError):
                    db.finish_document(row, {"status": "done", "text": "secret token",
                                             "method": "text", "fields": {"bad": {object()}}})
                state = db.conn.execute("SELECT parse_status,full_text FROM candidate_documents").fetchone()
                self.assertEqual((state[0], state[1]), ("running", ""))
                self.assertEqual(db.query(search="secret")[1], 0)
            finally:
                db.close()

    def test_changed_file_metadata_requeues_done_document(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(os.path.join(tmp, "test.db"), log=lambda *_: None).open()
            try:
                identity = {"source": "topcv", "account": "a", "cv_id": "1"}
                db.upsert({**identity, "filename": "1.txt", "dl_status": DONE})
                db.enqueue_document(**identity, filename="1.txt", file_size=10, file_mtime_ns=1)
                db.finish_document(db.claim_document(), {"status": "done", "text": "Python",
                    "method": "text", "quality": 1, "fields": {}, "file_size": 10,
                    "file_mtime_ns": 1})
                db.enqueue_document(**identity, filename="1.txt", file_size=11, file_mtime_ns=2)
                self.assertEqual(db.conn.execute(
                    "SELECT parse_status FROM candidate_documents").fetchone()[0], "pending")
            finally:
                db.close()

    def test_parsing_report_respects_candidate_filters(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(os.path.join(tmp, "test.db"), log=lambda *_: None).open()
            try:
                for source, cv_id in (("topcv", "1"), ("careerviet", "2")):
                    db.upsert({"source": source, "account": "a", "cv_id": cv_id,
                               "fullname": "Candidate " + cv_id, "filename": cv_id + ".txt",
                               "dl_status": DONE})
                    db.enqueue_document(source, "a", cv_id, cv_id + ".txt")
                row = db.claim_document()
                self.assertEqual(row["fullname"], "Candidate 1")
                db.finish_document(row, {"status": "done", "text": "Python",
                    "method": "text", "quality": 1, "fields": {}})
                report = db.stats_parsing(source=["topcv"])
                self.assertEqual(report["total"], 1)
                self.assertEqual(report["done"], 1)
                self.assertEqual(report["coverage"], 100.0)
                self.assertEqual(report["by_method"][0]["value"], "text")
            finally:
                db.close()

    def test_retry_failed_parsing_does_not_requeue_success(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(os.path.join(tmp, "test.db"), log=lambda *_: None).open()
            try:
                for cv_id in ("1", "2"):
                    db.upsert({"source": "topcv", "account": "a", "cv_id": cv_id,
                               "filename": cv_id + ".txt", "dl_status": DONE})
                    db.enqueue_document("topcv", "a", cv_id, cv_id + ".txt")
                    row = db.claim_document()
                    db.finish_document(row, {"status": "done" if cv_id == "1" else "error",
                        "text": "OK" if cv_id == "1" else "", "method": "text",
                        "quality": 1 if cv_id == "1" else 0, "fields": {}})
                self.assertEqual(db.requeue_failed_documents(), 1)
                states = {row["cv_id"]: row["parse_status"] for row in db.conn.execute(
                    "SELECT cv_id,parse_status FROM candidate_documents")}
                self.assertEqual(states, {"1": "done", "2": "pending"})
            finally:
                db.close()

    def test_parsing_queue_claim_and_retry_respect_priority_filters(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(os.path.join(tmp, "test.db"), log=lambda *_: None).open()
            try:
                candidates = [
                    ("topcv", "a@test.vn", "1", "Data Engineer", "2026-08-10 09:00:00"),
                    ("topcv", "b@test.vn", "2", "Data Engineer", "2026-08-10 09:00:00"),
                    ("careerviet", "a@test.vn", "3", "Data Engineer", "2026-08-10 09:00:00"),
                    ("topcv", "a@test.vn", "4", "Sales", "2026-07-01 09:00:00"),
                ]
                for source, account, cv_id, position, applied_ts in candidates:
                    db.upsert({"source": source, "account": account, "cv_id": cv_id,
                               "position": position, "applied_ts": applied_ts,
                               "filename": cv_id + ".txt", "dl_status": DONE})
                filters = {"source": ["topcv"], "account": ["a@test.vn"],
                           "position": ["Data Engineer"],
                           "date_from": "2026-08-01", "date_to": "2026-08-31"}
                page = db.enqueue_unparsed_batch(PARSER_VERSION, False, 0, 50, filters)
                self.assertEqual(page["count"], 1)
                self.assertEqual(db.claim_document(filters)["cv_id"], "1")

                # Một pending ngoài nhóm không được claim trong lượt ưu tiên hiện tại.
                db.enqueue_document("topcv", "b@test.vn", "2", "2.txt")
                self.assertIsNone(db.claim_document(filters))
                row = db.claim_document({"account": ["b@test.vn"]})
                db.finish_document(row, {"status": "error", "text": "", "fields": {}})
                self.assertEqual(db.requeue_failed_documents(filters), 0)
                self.assertEqual(db.requeue_failed_documents({"account": ["b@test.vn"]}), 1)
            finally:
                db.close()

    def test_parsing_priority_newest_to_oldest(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(os.path.join(tmp, "test.db"), log=lambda *_: None).open()
            try:
                candidates = [
                    ("topcv", "a@test.vn", "old_1", "2026-08-01 08:00:00"),
                    ("topcv", "a@test.vn", "mid_2", "2026-08-10 08:00:00"),
                    ("topcv", "a@test.vn", "new_3", "2026-08-17 14:00:00"),
                ]
                for source, account, cv_id, applied_ts in candidates:
                    db.upsert({"source": source, "account": account, "cv_id": cv_id,
                               "fullname": f"Candidate {cv_id}", "applied_ts": applied_ts,
                               "filename": cv_id + ".txt", "dl_status": DONE})
                    db.enqueue_document(source, account, cv_id, cv_id + ".txt")

                # Claim document phải ưu tiên hồ sơ mới nhất (new_3 -> mid_2 -> old_1)
                first = db.claim_document()
                self.assertEqual(first["cv_id"], "new_3")
                db.finish_document(first, {"status": "done", "text": "New text", "quality": 1})

                second = db.claim_document()
                self.assertEqual(second["cv_id"], "mid_2")
                db.finish_document(second, {"status": "done", "text": "Mid text", "quality": 1})

                third = db.claim_document()
                self.assertEqual(third["cv_id"], "old_1")
                db.finish_document(third, {"status": "done", "text": "Old text", "quality": 1})

                # Thêm hồ sơ pending_4 chưa từng qua parsing
                db.upsert({"source": "topcv", "account": "a@test.vn", "cv_id": "pending_4",
                           "fullname": "Candidate pending_4", "applied_ts": "2026-08-18 10:00:00",
                           "filename": "pending_4.txt", "dl_status": DONE})
                db.enqueue_document("topcv", "a@test.vn", "pending_4", "pending_4.txt")

                # Kiểm tra query_parsing_documents chỉ lấy 3 hồ sơ đã qua parsing, loại bỏ pending_4
                docs, total = db.query_parsing_documents(limit=10, offset=0)
                self.assertEqual(total, 3)
                self.assertEqual(len(docs), 3)
                self.assertTrue(all(d["cv_id"] != "pending_4" for d in docs))
            finally:
                db.close()

    def test_auto_fill_and_merge_contacts_from_parsing(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(os.path.join(tmp, "test.db"), log=lambda *_: None).open()
            try:
                # Ứng viên 1: Chưa có email & phone lúc tải về
                db.upsert({"source": "topcv", "account": "a", "cv_id": "cand_1",
                           "fullname": "Nguyen Van A", "email": "", "phone": "",
                           "filename": "cand_1.pdf", "dl_status": DONE})
                db.enqueue_document("topcv", "a", "cand_1", "cand_1.pdf")
                task1 = db.claim_document()
                db.finish_document(task1, {
                    "status": "done",
                    "text": "CV Nguyen Van A, email: nva@gmail.com, phone: 0912 345 678",
                    "quality": 1,
                    "fields": {
                        "emails": ["nva@gmail.com"],
                        "phones": ["0912345678"],
                        "urls": []
                    }
                })

                # Kiểm tra cand_1 được tự động điền email và phone
                cand1 = db.get_candidate("topcv", "cand_1", "a")
                self.assertEqual(cand1["email"], "nva@gmail.com")
                self.assertEqual(cand1["phone"], "0912345678")

                # Ứng viên 2: Đã có email và phone từ trước, parser phát hiện thêm email và phone thứ 2
                db.upsert({"source": "topcv", "account": "a", "cv_id": "cand_2",
                           "fullname": "Tran Thi B", "email": "main_b@company.com", "phone": "0988888888",
                           "filename": "cand_2.pdf", "dl_status": DONE})
                db.enqueue_document("topcv", "a", "cand_2", "cand_2.pdf")
                task2 = db.claim_document()
                db.finish_document(task2, {
                    "status": "done",
                    "text": "CV Tran Thi B, lien he: personal_b@gmail.com, 0977777777",
                    "quality": 1,
                    "fields": {
                        "emails": ["personal_b@gmail.com"],
                        "phones": ["0977777777"],
                        "urls": []
                    }
                })

                # `email`/`phone` giữ ĐÚNG MỘT giá trị — giá trị lấy từ trang nhà
                # tuyển dụng. Hub dùng ô này làm định danh Person, nên gộp nhiều
                # giá trị vào đây là trộn lẫn ứng viên với người tham chiếu ghi
                # trong CV; Hub khi đó không đọc nổi chuỗi có dấu phẩy và vứt
                # luôn cả định danh (đo thật: 28 hồ sơ mất email, 32 mất SĐT).
                cand2 = db.get_candidate("topcv", "cand_2", "a")
                self.assertEqual(cand2["email"], "main_b@company.com")
                self.assertEqual(cand2["phone"], "0988888888")

                # Giá trị bóc từ CV không mất — nằm riêng để Hub phân xử của ai.
                self.assertEqual(json.loads(cand2["cv_emails"]), ["personal_b@gmail.com"])
                self.assertEqual(json.loads(cand2["cv_phones"]), ["0977777777"])

                # Kiểm tra tìm kiếm FTS5 vẫn tìm được cả 2 email và cả 2 SĐT
                results, total = db.query(search="personal_b@gmail.com")
                self.assertEqual(total, 1)
                self.assertEqual(results[0]["cv_id"], "cand_2")

                results_phone, total_phone = db.query(search="0977777777")
                self.assertEqual(total_phone, 1)
                self.assertEqual(results_phone[0]["cv_id"], "cand_2")

                # Ứng viên 3 (ITViec - không có "nơi làm việc mong muốn" trên
                # trang chi tiết): lấy dự phòng từ text CV, chỉ điền khi trống.
                db.upsert({"source": "itviec", "account": "a", "cv_id": "cand_3",
                           "fullname": "Le Van C", "desired_location": "",
                           "filename": "cand_3.pdf", "dl_status": DONE})
                db.enqueue_document("itviec", "a", "cand_3", "cand_3.pdf")
                task3 = db.claim_document()
                db.finish_document(task3, {
                    "status": "done", "quality": 1,
                    "text": "Ho ten: Le Van C\nDia diem lam viec mong muon: Da Nang\nMuc tieu nghe nghiep: ...",
                    "fields": {"emails": [], "phones": [], "urls": [],
                               "desired_location": "Da Nang"},
                })
                cand3 = db.get_candidate("itviec", "cand_3", "a")
                self.assertEqual(cand3["desired_location"], "Da Nang")

                # Ứng viên 4: đã có desired_location từ trang chi tiết -> CV KHÔNG đè.
                db.upsert({"source": "vietnamworks", "account": "a", "cv_id": "cand_4",
                           "fullname": "Pham Thi D", "desired_location": "Ha Noi",
                           "filename": "cand_4.pdf", "dl_status": DONE})
                db.enqueue_document("vietnamworks", "a", "cand_4", "cand_4.pdf")
                task4 = db.claim_document()
                db.finish_document(task4, {
                    "status": "done", "quality": 1,
                    "text": "Noi lam viec mong muon: TP HCM",
                    "fields": {"emails": [], "phones": [], "urls": [],
                               "desired_location": "TP HCM"},
                })
                self.assertEqual(db.get_candidate("vietnamworks", "cand_4", "a")["desired_location"], "Ha Noi")
            finally:
                db.close()

    def test_email_nguoi_tham_chieu_khong_bi_gan_cho_ung_vien(self):
        """CV hay kèm liên hệ NGƯỜI THAM CHIẾU. Hồ sơ thật 'Nguyễn Hải Yến' có
        thêm email hai banker SeABank/VPBank. Ô `email` (Hub dùng làm định danh
        Person) phải là của chính ứng viên, không phải của người tham chiếu."""
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(os.path.join(tmp, "ref.db"), log=lambda *_: None).open()
            try:
                db.upsert({"source": "topcv", "account": "a", "cv_id": "yen",
                           "fullname": "Nguyễn Hải Yến", "email": "", "phone": "",
                           "filename": "yen.pdf", "dl_status": DONE})
                db.enqueue_document("topcv", "a", "yen", "yen.pdf")
                db.finish_document(db.claim_document(), {
                    "status": "done", "quality": 1,
                    "text": "NGUOI THAM CHIEU: Thuy PT - SeABank",
                    "fields": {"urls": [],
                               "emails": ["nguyenhaiyen105@gmail.com",
                                          "thuy.pt@seabank.com.vn",
                                          "phuongntm39@vpbank.com.vn"],
                               "phones": ["0972748510", "0986672829"]},
                })
                row = db.get_candidate("topcv", "yen", "a")
                self.assertEqual(row["email"], "nguyenhaiyen105@gmail.com")
                self.assertEqual(row["phone"], "0972748510")
                # Không mất gì: cả ba vẫn nằm ở cột riêng cho Hub phân xử.
                self.assertEqual(len(json.loads(row["cv_emails"])), 3)
                self.assertIn("thuy.pt@seabank.com.vn", json.loads(row["cv_emails"]))
            finally:
                db.close()

    def test_migration_tach_lai_o_lien_he_bi_gop_cua_ban_cu(self):
        """Database phiên bản cũ đã có sẵn ô gộp 'a@x, b@y'. Lần mở đầu tiên sau
        khi nâng cấp phải tách lại, nếu không Hub vẫn vứt định danh như trước."""
        import sqlite3
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "legacy.db")
            db = Database(path, log=lambda *_: None).open()
            try:
                db.upsert({"source": "topcv", "account": "a", "cv_id": "old",
                           "fullname": "Nguyễn Hải Yến", "dl_status": DONE})
            finally:
                db.close()
            # Dựng lại đúng trạng thái bản cũ: ô gộp + chưa có 2 cột mới.
            raw = sqlite3.connect(path)
            raw.execute("UPDATE candidates SET email=?, phone=?",
                        ("nguyenhaiyen105@gmail.com, thuy.pt@seabank.com.vn",
                         "0972748510, 0986672829"))
            raw.execute("ALTER TABLE candidates DROP COLUMN cv_emails")
            raw.execute("ALTER TABLE candidates DROP COLUMN cv_phones")
            raw.commit()
            raw.close()

            db = Database(path, log=lambda *_: None).open()
            try:
                row = db.get_candidate("topcv", "old", "a")
                self.assertEqual(row["email"], "nguyenhaiyen105@gmail.com")
                self.assertEqual(row["phone"], "0972748510")
                self.assertEqual(json.loads(row["cv_emails"]),
                                 ["nguyenhaiyen105@gmail.com", "thuy.pt@seabank.com.vn"])
            finally:
                db.close()

    def test_tach_lai_o_gop_ke_ca_khi_da_co_san_cot_moi(self):
        """Cột `cv_emails` đã tồn tại nhưng dữ liệu vẫn còn gộp — xảy ra thật khi
        DB đi qua nhánh dựng lại bảng của `_migrate_account_key` (bảng mới có đủ
        cột rồi mới chép dữ liệu cũ vào). Nếu việc tách gài vào nhánh "vừa thêm
        cột" thì không nhánh nào chạy và dữ liệu gộp sống sót vĩnh viễn."""
        import sqlite3
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "sau_rebuild.db")
            db = Database(path, log=lambda *_: None).open()
            try:
                db.upsert({"source": "topcv", "account": "a", "cv_id": "x",
                           "fullname": "Nguyễn Hải Yến", "dl_status": DONE})
            finally:
                db.close()
            raw = sqlite3.connect(path)   # cột mới CÒN NGUYÊN, chỉ dữ liệu bị gộp
            raw.execute("UPDATE candidates SET email=?, phone=?, cv_emails=NULL",
                        ("a@gmail.com, b@seabank.com.vn", "0972748510, 0986672829"))
            raw.commit()
            raw.close()

            db = Database(path, log=lambda *_: None).open()
            try:
                row = db.get_candidate("topcv", "x", "a")
                self.assertNotIn(",", row["email"])
                self.assertNotIn(",", row["phone"])
                self.assertEqual(len(json.loads(row["cv_emails"])), 2)
            finally:
                db.close()

    def test_ascii_path_copies_vietnamese_filename_for_subprocess_tools(self):
        """Tesseract/LibreOffice trên Windows làm hỏng đường dẫn có dấu tiếng
        Việt ("Lê Hương Dậu.png" -> "...D?u.png" -> "cannot read input file").
        `_ascii_path` chép sang tên ASCII rồi xoá sau khi xong."""
        from app.cv_parser import _ascii_path
        with tempfile.TemporaryDirectory() as tmp:
            with _ascii_path(os.path.join(tmp, "ok_ascii.png")) as p:
                self.assertTrue(p.endswith("ok_ascii.png"))  # ascii -> trả nguyên

            src = os.path.join(tmp, "4059463_Lê Hương Dậu.png")
            with open(src, "wb") as handle:
                handle.write(b"PNGDATA")
            with _ascii_path(src) as p:
                p.encode("ascii")               # không ném lỗi
                self.assertTrue(os.path.isfile(p))
                with open(p, "rb") as handle:
                    self.assertEqual(handle.read(), b"PNGDATA")
                copied = p
            self.assertFalse(os.path.exists(copied))  # đã dọn

    def test_desired_location_from_cv_reads_common_vietnamese_labels(self):
        from app.cv_parser import _basic_fields
        self.assertEqual(
            _basic_fields("... Địa điểm làm việc mong muốn: Đà Nẵng\nMục tiêu nghề nghiệp")["desired_location"],
            "Đà Nẵng")
        self.assertEqual(
            _basic_fields("Nơi làm việc mong muốn : Hồ Chí Minh, Bình Dương . Kinh nghiệm")["desired_location"],
            "Hồ Chí Minh, Bình Dương")
        self.assertEqual(_basic_fields("CV không ghi nơi mong muốn gì cả")["desired_location"], "")

    def test_docx_textbox_content_is_extracted(self):
        """CV Word đặt nội dung trong text box: python-docx trả rỗng, phải quét
        thô XML để lấy được."""
        marker = "KY NANG Python SQL " + ("noi dung trong text box " * 8)
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "textbox_cv.docx")
            _docx_with_textbox(path, "   ", marker)

            from docx import Document
            native = "\n".join(p.text for p in Document(path).paragraphs).strip()
            self.assertEqual(native, "", "tiền đề: python-docx không thấy text box")

            result = extract_cv(path)
            self.assertEqual(result["status"], "done")
            self.assertIn("noi dung trong text box", result["text"])
            self.assertEqual(result["method"], "docx+xml")

    def test_docx_normal_still_uses_python_docx(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "plain.docx")
            from docx import Document
            document = Document()
            document.add_paragraph(
                "Nguyen Van A - Ky su du lieu. Kinh nghiem 5 nam voi Python, SQL, "
                "Apache Airflow, dbt va Spark. Tung xay dung pipeline ETL cho ngan "
                "hang va cong ty fintech. Tot nghiep Bach Khoa Ha Noi nam 2018.")
            document.save(path)
            result = extract_cv(path)
            self.assertEqual(result["status"], "done")
            self.assertEqual(result["method"], "docx")
            self.assertIn("Airflow", result["text"])

    def test_parsing_backlog_and_stale_retry_helpers(self):
        with tempfile.TemporaryDirectory() as tmp:
            db = Database(os.path.join(tmp, "test.db"), log=lambda *_: None).open()
            try:
                for cv_id in ("1", "2", "3"):
                    db.upsert({"source": "topcv", "account": "a", "cv_id": cv_id,
                               "filename": cv_id + ".txt", "dl_status": DONE})
                # 3 CV đã tải, chưa CV nào parse -> đều là tồn đọng.
                self.assertEqual(db.parsing_backlog(PARSER_VERSION), 3)
                self.assertEqual(db.count_parsed_terminal(), 0)

                db.enqueue_document("topcv", "a", "1", "1.txt")
                db.finish_document(db.claim_document(), {"status": "done", "text": "ok",
                    "method": "text", "quality": 1, "fields": {}})
                # CV 1 xong -> tồn đọng còn 2 (CV 2, 3 chưa có bản ghi tài liệu).
                self.assertEqual(db.parsing_backlog(PARSER_VERSION), 2)
                self.assertEqual(db.count_parsed_terminal(), 1)

                db.enqueue_document("topcv", "a", "2", "2.txt")
                db.finish_document(db.claim_document(), {"status": "error", "text": "",
                    "method": "", "quality": 0, "fields": {}})
                # CV 2 lỗi: KHÔNG tính vào tồn đọng thường (chờ nguội), tính terminal.
                self.assertEqual(db.parsing_backlog(PARSER_VERSION), 1)
                self.assertEqual(db.count_parsed_terminal(), 2)

                # Vừa lỗi xong -> chưa đủ nguội, không xếp lại.
                self.assertEqual(db.requeue_stale_failed_documents(3, 6), 0)
                # Ép updated_at về quá khứ để mô phỏng đã nguội.
                db.conn.execute("UPDATE candidate_documents SET updated_at='2000-01-01 00:00:00' "
                                "WHERE cv_id='2'")
                self.assertEqual(db.parsing_retry_backlog(3, 6), 1)
                self.assertEqual(db.requeue_stale_failed_documents(3, 6), 1)
                self.assertEqual(db.conn.execute(
                    "SELECT parse_status FROM candidate_documents WHERE cv_id='2'"
                ).fetchone()[0], "pending")
            finally:
                db.close()


if __name__ == "__main__":
    unittest.main()
