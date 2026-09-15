# -*- coding: utf-8 -*-
"""Xuất dữ liệu ứng viên ra Excel (.xlsx) hoặc CSV để gửi cho người khác."""
import csv
import os
import zipfile

from .db import DONE, EXPORT_COLUMNS


def export_rows(rows, path, fmt=None, cv_folder=None):
    """
    rows      : danh sách bản ghi lấy từ Database.query_all()
    fmt       : 'xlsx' | 'csv' (để trống thì suy ra từ đuôi file)
    cv_folder : nếu có VÀ fmt là xlsx, cột "Tên file CV" sẽ thành liên kết bấm-mở-được
                tới đúng file CV trên đĩa (chỉ với ứng viên đã tải xong và file còn tồn tại).
                CSV không hỗ trợ liên kết thật nên bỏ qua tham số này.
    Trả về số dòng đã xuất.
    """
    fmt = (fmt or os.path.splitext(path)[1].lstrip(".")).lower()
    folder = os.path.dirname(path)
    if folder:
        os.makedirs(folder, exist_ok=True)
    if fmt == "csv":
        return _to_csv(rows, path)
    return _to_xlsx(rows, path, cv_folder=cv_folder)


def _local_cv_links(rows, cv_folder):
    """Đường dẫn tuyệt đối tới file CV trên đĩa cho từng dòng (None nếu chưa tải/không thấy file)."""
    out = []
    for r in rows:
        fn = _value(r, "filename")
        if r["dl_status"] == DONE and fn:
            p = os.path.join(cv_folder, str(fn))
            out.append(p if os.path.exists(p) else None)
        else:
            out.append(None)
    return out


def _value(r, key):
    try:
        v = r[key]
    except (KeyError, IndexError):
        return ""
    if key == "is_viewed":
        return "Có" if v else "Chưa"
    return "" if v is None else v


def _to_csv(rows, path):
    # utf-8-sig để Excel trên Windows mở tiếng Việt không bị lỗi phông
    count = 0
    with open(path, "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f)
        w.writerow([label for _, label in EXPORT_COLUMNS])
        for r in rows:
            w.writerow([_value(r, key) for key, _ in EXPORT_COLUMNS])
            count += 1
    return count


_FILENAME_COL = next(i for i, (key, _) in enumerate(EXPORT_COLUMNS, start=1) if key == "filename")


def _to_xlsx(rows, path, links=None, cv_folder=None):
    """
    links: (tuỳ chọn) danh sách song song với rows - links[i] là đường dẫn/liên kết
           tới file CV của dòng thứ i (tuyệt đối hoặc tương đối), hoặc None nếu dòng
           đó không có file kèm theo. Có giá trị thì cột "Tên file CV" sẽ bấm mở được.
    """
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook(write_only=True)
    ws = wb.create_sheet()
    ws.title = "Ứng viên"
    ws.freeze_panes = "A2"

    fill = PatternFill("solid", fgColor="1F6F4A")
    font = Font(color="FFFFFF", bold=True)
    header = []
    for _, label in EXPORT_COLUMNS:
        c = WriteOnlyCell(ws, value=label)
        c.fill, c.font = fill, font
        c.alignment = Alignment(vertical="center")
        header.append(c)
    ws.append(header)

    # Bề rộng theo KHÓA cột (không theo vị trí) để chèn thêm cột ở giữa
    # EXPORT_COLUMNS không làm lệch bề rộng của các cột khác.
    width_by_key = {
        "source": 10, "account": 22, "cv_id": 16, "fullname": 24, "email": 26,
        "phone": 15, "position": 30, "campaign_id": 14, "applied_at": 16,
        "apply_source": 22, "status": 14, "gender": 10, "birth_year": 10,
        "marital_status": 16, "experience": 16, "years_experience": 14,
        "address": 30, "city": 16, "district": 16, "desired_location": 22,
        "current_title": 22, "job_level": 14, "desired_level": 16,
        "desired_position": 24, "job_type": 24, "education": 18,
        "foreign_language": 18, "expected_salary": 16, "current_salary": 16,
        "skills": 32, "last_company": 24, "labels": 20, "note": 28,
        "candidate_id": 16, "resume_id": 16, "profile_type": 12,
        "attachment_name": 24, "attachment_mime": 18, "is_viewed": 9,
        "filename": 32, "dl_status": 16, "alerts": 40, "cv_url": 44,
        "first_seen": 19, "updated_at": 19,
    }
    for i, (key, _label) in enumerate(EXPORT_COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width_by_key.get(key, 16)

    link_font = Font(color="0563C1", underline="single")
    count = 0
    for row_i, r in enumerate(rows, start=2):
        values = [_value(r, key) for key, _ in EXPORT_COLUMNS]
        target = links[row_i - 2] if links else None
        if not target and cv_folder and _value(r, "dl_status") == DONE:
            filename = _value(r, "filename")
            candidate = os.path.join(cv_folder, str(filename)) if filename else ""
            target = candidate if candidate and os.path.exists(candidate) else None
        if target:
            c = WriteOnlyCell(ws, value=values[_FILENAME_COL - 1])
            c.hyperlink = target
            c.font = link_font
            values[_FILENAME_COL - 1] = c
        ws.append(values)
        count += 1

    ws.auto_filter.ref = f"A1:{get_column_letter(len(EXPORT_COLUMNS))}{max(1, count + 1)}"
    wb.save(path)
    return count


def _dedupe_filename(name, used_names):
    """Tránh 2 file trùng tên trong cùng 1 file nén (hiếm gặp vì tên file đã kèm mã CV)."""
    if name not in used_names:
        used_names.add(name)
        return name
    base, ext = os.path.splitext(name)
    i = 2
    while f"{base}_{i}{ext}" in used_names:
        i += 1
    new_name = f"{base}_{i}{ext}"
    used_names.add(new_name)
    return new_name


def export_zip_with_cv(rows, zip_path, cv_folder, on_progress=None):
    """
    Xuất danh sách Excel + toàn bộ file CV tương ứng của kết quả lọc vào 1 file .zip,
    để gửi cho người khác mà không cần đi tìm từng file CV thủ công.

    rows      : danh sách bản ghi lấy từ Database.query_all(**bộ_lọc_hiện_tại)
    cv_folder : thư mục chứa các file CV đã tải (AppConfig.cv_folder)
    on_progress(i, total): gọi sau mỗi ứng viên xử lý xong, có thể để trống.

    Chỉ ứng viên đã tải xong (dl_status == DONE) VÀ file còn tồn tại trên đĩa mới được
    kèm theo - các trường hợp còn lại (chưa tải / file bị xoá, di chuyển) vẫn có mặt đầy
    đủ trong Excel như bình thường, nhưng được liệt kê riêng trong "Danh_sach_thieu_CV.txt"
    bên trong file nén để người nhận biết vì sao thiếu, không phải đoán. Cột "Tên file CV"
    trong Excel bấm được thẳng vào file CV tương ứng nằm cạnh nó trong file nén.
    """
    folder = os.path.dirname(zip_path)
    if folder:
        os.makedirs(folder, exist_ok=True)

    total = len(rows)
    not_downloaded = missing_file = 0
    missing_names = []
    used_names = set()
    # đường dẫn trong file nén (đường / theo chuẩn zip) cho từng dòng, None = không có file
    arcnames = [None] * total
    for i, r in enumerate(rows):
        fn = _value(r, "filename")
        done = r["dl_status"] == DONE
        name = _value(r, "fullname") or _value(r, "cv_id")
        if not done or not fn:
            not_downloaded += 1
            if name and len(missing_names) < 200:
                missing_names.append(f"{name} - chưa tải CV")
            continue
        src = os.path.join(cv_folder, str(fn))
        if os.path.exists(src):
            arcnames[i] = "CV/" + _dedupe_filename(str(fn), used_names)
        else:
            missing_file += 1
            if name and len(missing_names) < 200:
                missing_names.append(f"{name} - không tìm thấy file trên đĩa")

    tmp_xlsx = zip_path + ".tmp.xlsx"
    try:
        # liên kết dùng đường \ để Excel/Windows hiểu là đường dẫn TƯƠNG ĐỐI kể từ nơi
        # đặt file Excel này - đúng vì sau khi giải nén, "CV/" luôn nằm cạnh file Excel.
        links = [a.replace("/", "\\") if a else None for a in arcnames]
        _to_xlsx(rows, tmp_xlsx, links=links)

        included = 0
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            zf.write(tmp_xlsx, "Danh_sach_ung_vien.xlsx")
            for i, r in enumerate(rows, start=1):
                arcname = arcnames[i - 1]
                if arcname:
                    src = os.path.join(cv_folder, str(_value(r, "filename")))
                    zf.write(src, arcname)
                    included += 1
                if on_progress:
                    on_progress(i, total)

            if missing_names:
                report = ["Các ứng viên KHÔNG có file CV kèm theo trong lần xuất này:", ""]
                report += missing_names
                left = (not_downloaded + missing_file) - len(missing_names)
                if left > 0:
                    report.append(f"... và {left:,} người khác.")
                zf.writestr("Danh_sach_thieu_CV.txt", "\n".join(report))
    finally:
        if os.path.exists(tmp_xlsx):
            os.remove(tmp_xlsx)

    return {
        "total": total, "included": included,
        "not_downloaded": not_downloaded, "missing_file": missing_file,
    }
