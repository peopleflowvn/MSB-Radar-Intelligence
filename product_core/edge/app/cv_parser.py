# -*- coding: utf-8 -*-
"""Bounded, durable and resumable CV text extraction pipeline."""
import contextlib
import hashlib
import os
import re
import shutil
import threading
import time
import subprocess
import tempfile
import zipfile
import xml.etree.ElementTree as ET
from collections import deque
from datetime import datetime
from pathlib import Path

PARSER_VERSION = "2"
IMAGE_SUFFIXES = {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp", ".webp"}
LEGACY_OFFICE = {".doc": ".docx", ".xls": ".xlsx", ".ppt": ".pptx"}


@contextlib.contextmanager
def _ascii_path(path):
    """Trả một đường dẫn CHỈ CÓ ASCII cho `path`.

    Tesseract / LibreOffice trên Windows nhận đường dẫn có dấu tiếng Việt qua
    argv sẽ chuyển sang codepage ANSI và làm hỏng ký tự ("Lê Hương Dậu.png" ->
    "Lê Huong D?u.png" -> "cannot read input file ... Invalid argument"). Nếu
    đường dẫn có ký tự ngoài ASCII thì chép sang file tạm tên ASCII, xong thì
    xoá. Đường dẫn thuần ASCII trả nguyên, không chép gì."""
    try:
        str(path).encode("ascii")
        yield path
        return
    except UnicodeEncodeError:
        pass
    handle, temp = tempfile.mkstemp(prefix="msbradar_ascii_", suffix=Path(path).suffix.lower())
    os.close(handle)
    try:
        shutil.copyfile(path, temp)
        yield temp
    finally:
        try:
            os.remove(temp)
        except OSError:
            pass


def _clean(text):
    text = str(text or "").replace("\x00", " ")
    lines = [re.sub(r"[ \t]+", " ", line).strip() for line in text.splitlines()]
    return "\n".join(line for line in lines if line).strip()


def _hash_file(path):
    digest = hashlib.sha256()
    with open(path, "rb") as stream:
        for block in iter(lambda: stream.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def _normalize_phone(p):
    cleaned = re.sub(r"[^\d+]", "", str(p or "")).strip()
    if cleaned.startswith("+84"):
        cleaned = "0" + cleaned[3:]
    elif cleaned.startswith("84") and len(cleaned) == 11:
        cleaned = "0" + cleaned[2:]
    return cleaned


def _basic_fields(text):
    text_str = str(text or "")
    # Trích xuất Email chính xác, lọc đuôi file ảnh hoặc ký tự thừa
    raw_emails = re.findall(r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+", text_str)
    valid_emails = []
    seen_emails = set()
    for em in raw_emails:
        em_clean = em.strip(".,;:()[]{}<>\"' \t\n\r").lower()
        if "@" in em_clean and "." in em_clean.split("@")[-1] and not em_clean.endswith((".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg")):
            if em_clean not in seen_emails:
                seen_emails.add(em_clean)
                valid_emails.append(em_clean)

    # Trích xuất Số điện thoại chuẩn định dạng VN
    raw_phones = re.findall(r"(?<!\d)(?:\+?84|0)[\d .()-]{7,15}\d", text_str)
    valid_phones = []
    seen_phones = set()
    for ph in raw_phones:
        norm = _normalize_phone(ph)
        if norm.startswith("0") and 9 <= len(norm) <= 11:
            if norm not in seen_phones:
                seen_phones.add(norm)
                valid_phones.append(norm)

    urls = sorted(set(re.findall(r"https?://[^\s<>'\"]+", text_str, re.I)))[:20]
    return {"emails": valid_emails[:10], "phones": valid_phones[:10], "urls": urls,
            "desired_location": _desired_location_from_cv(text_str)}


#: Nhãn "nơi làm việc mong muốn" hay gặp trong CV tiếng Việt (mỗi mẫu CV một
#: kiểu). PHƯƠNG ÁN DỰ PHÒNG khi trang chi tiết của nguồn không cho được trường
#: này (ITViec/Joboko, hoặc khi bóc web thất bại). Ưu tiên vẫn là giá trị từ
#: web (xem `db.finish_document_v2` chỉ điền khi cột trống). Xem
#: KINH_NGHIEM_TRIEN_KHAI_CAC_KENH_TUYEN_DUNG.md mục 41.
_DESIRED_LOCATION_RE = re.compile(
    r"(?:địa\s*điểm|khu\s*vực)\s*(?:làm\s*việc\s*)?mong\s*muốn"
    r"|nơi\s*làm\s*việc\s*mong\s*muốn"
    r"|nơi\s*mong\s*muốn\s*làm\s*việc"
    r"|(?:desired|preferred|expected)\s*(?:work\s*)?location",
    re.I)


def _desired_location_from_cv(text):
    """Địa điểm mong muốn suy từ text CV, theo 2 bước:

    1. Có nhãn "địa điểm/nơi làm việc mong muốn" -> lấy chuỗi ngay sau, chuẩn
       hoá về tên tỉnh/thành (`geo.normalize_location`).
    2. Không có nhãn -> quét cả text tìm địa danh cấp tỉnh/thành (mọi cách viết,
       viết tắt) và lấy 2 cái đầu. Kém tin cậy hơn (có thể trúng nơi ở hiện tại)
       nhưng vẫn là thông tin địa điểm hữu ích khi không có gì khác.
    """
    from .geo import find_provinces, normalize_location

    text = text or ""
    match = _DESIRED_LOCATION_RE.search(text)
    if match:
        tail = text[match.end():match.end() + 120]
        tail = re.split(r"[\r\n]|(?:[.;•|]\s)|\s{3,}", tail.lstrip(" :：-–\t"), 1)[0]
        tail = tail.strip(" :：-–,\t")
        tail = re.split(r"\s+(?:mục tiêu|kinh nghiệm|học vấn|kỹ năng|objective|experience)\b",
                        tail, 1, re.I)[0].strip()
        if 1 <= len(tail) <= 120:
            return normalize_location(tail)

    provinces = find_provinces(text)
    return ", ".join(provinces[:2]) if provinces else ""


def _tesseract_settings(languages="vie+eng"):
    from .prerequisites import tesseract_runtime
    executable, local_data = tesseract_runtime()
    if not executable:
        raise RuntimeError("Tesseract OCR chưa sẵn sàng")
    env = os.environ.copy()
    if local_data and os.path.isfile(os.path.join(local_data, "eng.traineddata")):
        env["TESSDATA_PREFIX"] = local_data
    try:
        completed = subprocess.run([executable, "--list-langs"], capture_output=True,
                                   text=True, timeout=15, env=env,
                                   creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
        available = set(completed.stdout.splitlines()[1:])
    except Exception as exc:
        raise RuntimeError(f"Tesseract OCR chưa sẵn sàng: {exc}") from exc
    selected = [lang for lang in languages.split("+") if lang in available]
    if not selected:
        raise RuntimeError("Tesseract chưa có gói ngôn ngữ vie hoặc eng")
    return executable, env, "+".join(selected)


def _ocr_image(path, languages="vie+eng"):
    executable, env, selected = _tesseract_settings(languages)
    with _ascii_path(path) as safe_path:
        completed = subprocess.run([executable, safe_path, "stdout", "-l", selected],
                                   capture_output=True, text=True, timeout=90, env=env,
                                   creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
    if completed.returncode:
        raise RuntimeError((completed.stderr or "Tesseract OCR thất bại").strip())
    return completed.stdout


def _docx_xml_sweep(path):
    """Quét thô mọi ``<w:t>`` trong document.xml, header/footer và footnotes.

    python-docx chỉ đọc ``paragraphs`` và ``tables`` nên bỏ sót nội dung nằm trong
    text box (``w:txbxContent``), shape, WordArt và khối ``mc:AlternateContent`` -
    nhiều CV thiết kế đặt toàn bộ thông tin trong text box, và với các file đó
    python-docx trả về gần như rỗng. Các ``<w:t>`` này vẫn là con cháu trong cây
    XML nên chỉ cần duyệt phẳng toàn bộ phần thân là lấy được.
    """
    parts = []
    with zipfile.ZipFile(path) as archive:
        names = [name for name in archive.namelist()
                 if re.fullmatch(r"word/(document|header\d+|footer\d+|footnotes|endnotes)\.xml", name)]
        for name in sorted(names):
            try:
                root = ET.fromstring(archive.read(name))
            except ET.ParseError:
                continue
            for node in root.iter():
                tag = node.tag
                if tag.endswith("}t"):
                    parts.append(node.text or "")
                elif tag.endswith("}tab"):
                    parts.append("\t")
                elif tag.endswith("}br") or tag.endswith("}cr") or tag.endswith("}p"):
                    parts.append("\n")
    return "".join(parts)


def extract_cv(path, ocr_enabled=True, min_native_chars=120, max_pages=30,
               max_file_mb=50, max_sheet_cells=200000):
    suffix = Path(path).suffix.lower()
    text, method, needs_ocr, warning = "", "", False, ""
    stat = os.stat(path)
    if stat.st_size > max(1, int(max_file_mb)) * 1024 * 1024:
        return _result(path, suffix, "error", error=f"File vượt giới hạn {max_file_mb} MB")
    try:
        if suffix in LEGACY_OFFICE:
            from .prerequisites import libreoffice_runtime
            executable = libreoffice_runtime()
            if not executable:
                return _result(path, suffix, "unsupported",
                               error="Cần cài LibreOffice để đọc định dạng Office cũ")
            target_ext = LEGACY_OFFICE[suffix]
            with tempfile.TemporaryDirectory(prefix="msbradar_cv_") as folder, \
                    _ascii_path(path) as safe_path:
                conversion = {".doc": "docx", ".xls": "xlsx", ".ppt": "pptx"}[suffix]
                completed = subprocess.run(
                    [executable, "--headless", "--convert-to", conversion, "--outdir", folder,
                     safe_path],
                    capture_output=True, text=True, timeout=90,
                    creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
                converted = os.path.join(folder, Path(safe_path).stem + target_ext)
                if completed.returncode or not os.path.isfile(converted):
                    detail = (completed.stderr or completed.stdout or "không tạo được file").strip()
                    return _result(path, suffix, "error", error=f"LibreOffice: {detail}")
                result = extract_cv(converted, ocr_enabled, min_native_chars, max_pages,
                                    max_file_mb, max_sheet_cells)
                original = os.stat(path)
                result.update({"file_hash": _hash_file(path), "file_size": original.st_size,
                               "file_mtime_ns": original.st_mtime_ns,
                               "file_format": suffix.lstrip('.'),
                               "method": "libreoffice-" + result.get("method", conversion)})
                return result
        if suffix == ".pdf":
            import fitz
            with fitz.open(path) as doc:
                pages = [doc[index] for index in range(min(len(doc), max(1, int(max_pages))))]
                text = "\n".join(page.get_text("text") for page in pages)
                method = "pdf-native"
                needs_ocr = len(_clean(text)) < int(min_native_chars)
                if needs_ocr and ocr_enabled:
                    try:
                        chunks = []
                        with tempfile.TemporaryDirectory(prefix="msbradar_ocr_") as folder:
                            for index, page in enumerate(pages):
                                pix = page.get_pixmap(matrix=fitz.Matrix(1.7, 1.7), alpha=False)
                                image_path = os.path.join(folder, f"page_{index}.png")
                                pix.save(image_path)
                                chunks.append(_ocr_image(image_path))
                        ocr_text = _clean("\n".join(chunks))
                        if ocr_text:
                            text, method, needs_ocr = ocr_text, "pdf-ocr", False
                    except Exception as exc:
                        warning = str(exc)
        elif suffix in IMAGE_SUFFIXES:
            if not ocr_enabled:
                return _result(path, suffix, "needs_ocr", needs_ocr=True,
                               error="OCR đang tắt cho file ảnh")
            text = _ocr_image(path)
            method = "image-ocr"
        elif suffix == ".docx":
            from docx import Document
            doc = Document(path)
            parts = [p.text for p in doc.paragraphs]
            parts.extend(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
            native = "\n".join(parts)
            try:
                swept = _docx_xml_sweep(path)
            except Exception:
                swept = ""
            # document.xml chứa mọi <w:t>, kể cả trong text box - nên bản quét thô
            # là tập cha của những gì python-docx đọc được. Chỉ dùng nó khi thật sự
            # dài hơn (file CV toàn text box), còn lại giữ python-docx cho ổn định.
            if len(_clean(swept)) > len(_clean(native)):
                text, method = swept, "docx+xml"
            else:
                text, method = native, "docx"
            if len(_clean(text)) < int(min_native_chars):
                from .prerequisites import libreoffice_runtime
                executable = libreoffice_runtime()
                if executable:
                    try:
                        with tempfile.TemporaryDirectory(prefix="msbradar_cv_") as folder, \
                                _ascii_path(path) as safe_path:
                            completed = subprocess.run(
                                [executable, "--headless", "--convert-to", "pdf",
                                 "--outdir", folder, safe_path],
                                capture_output=True, text=True, timeout=90,
                                creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
                            converted = os.path.join(folder, Path(safe_path).stem + ".pdf")
                            if not completed.returncode and os.path.isfile(converted):
                                result = extract_cv(converted, ocr_enabled, min_native_chars,
                                                    max_pages, max_file_mb, max_sheet_cells)
                                original = os.stat(path)
                                result.update({"file_hash": _hash_file(path),
                                               "file_size": original.st_size,
                                               "file_mtime_ns": original.st_mtime_ns,
                                               "file_format": "docx",
                                               "method": "libreoffice-" + result.get("method", "pdf")})
                                return result
                    except Exception as exc:
                        warning = f"LibreOffice fallback thất bại: {exc}"
        elif suffix == ".pptx":
            with zipfile.ZipFile(path) as archive:
                slide_names = sorted(name for name in archive.namelist()
                                     if re.fullmatch(r"ppt/slides/slide\d+\.xml", name))
                parts = []
                for name in slide_names:
                    root = ET.fromstring(archive.read(name))
                    parts.extend(node.text or "" for node in root.iter()
                                 if node.tag.endswith("}t"))
            text = "\n".join(parts)
            method = "pptx"
        elif suffix in (".xlsx", ".xlsm"):
            from openpyxl import load_workbook
            book = load_workbook(path, read_only=True, data_only=True)
            lines, cells = [], 0
            for sheet in book.worksheets:
                for row in sheet.iter_rows(values_only=True):
                    cells += len(row)
                    if cells > max(1, int(max_sheet_cells)):
                        warning = f"Đã giới hạn ở {max_sheet_cells:,} ô"
                        break
                    lines.append(" | ".join(str(value) for value in row if value is not None))
                if warning:
                    break
            book.close()
            text, method = "\n".join(lines), "xlsx"
        elif suffix in (".txt", ".csv", ".md"):
            raw = Path(path).read_bytes()
            for encoding in ("utf-8-sig", "utf-16", "cp1258", "latin-1"):
                try:
                    text = raw.decode(encoding)
                    break
                except UnicodeError:
                    continue
            method = "text"
        elif suffix == ".rtf":
            raw = Path(path).read_bytes().decode("latin-1", errors="replace")
            text = re.sub(r"\\'[0-9a-fA-F]{2}", " ", raw)
            text = re.sub(r"\\[a-zA-Z]+-?\d* ?", " ", text)
            text = text.translate(str.maketrans({"{": " ", "}": " ", "\\": " "}))
            method = "rtf"
        else:
            return _result(path, suffix, "unsupported",
                           error=f"Định dạng {suffix or 'không xác định'} chưa được hỗ trợ")
        text = _clean(text)
        if needs_ocr:
            status = "needs_ocr"
        elif not text:
            status = "empty"
        else:
            status = "done"
        return _result(path, suffix, status, text, method, needs_ocr, warning)
    except Exception as exc:
        return _result(path, suffix, "error", method=method, needs_ocr=needs_ocr,
                       error=f"{type(exc).__name__}: {exc}")


def _result(path, suffix, status, text="", method="", needs_ocr=False, error=""):
    try:
        stat = os.stat(path)
        file_hash = _hash_file(path)
        size, mtime = stat.st_size, stat.st_mtime_ns
    except OSError:
        file_hash, size, mtime = "", 0, 0
    clean = _clean(text)
    useful = len(re.findall(r"\w+", clean, re.UNICODE))
    quality = min(1.0, useful / 180.0) if clean else 0
    return {"status": status, "text": clean, "method": method, "needs_ocr": needs_ocr,
            "quality": quality, "file_format": suffix.lstrip("."), "file_hash": file_hash,
            "file_size": size, "file_mtime_ns": mtime, "error": str(error)[:500],
            "fields": _basic_fields(clean), "parser_version": PARSER_VERSION}


class CvParsingService:
    def __init__(self, config, db_factory, log=print, on_finish=None):
        self.config, self.db_factory, self.log = config, db_factory, log
        self._on_finish = on_finish or (lambda: None)
        self._thread, self._stop, self._lock = None, threading.Event(), threading.Lock()
        self.status = {"processed": 0, "session_done": 0, "session_failed": 0}
        self._events = deque(maxlen=2000)
        self._event_seq = 0
        #: khác rỗng = đang tạm dừng nhường máy (tải/parsing ưu tiên hơn); hàng đợi giữ nguyên.
        self._paused_reason = ""
        #: người dùng chủ động bấm Dừng -> điều phối KHÔNG được tự chạy lại cho tới khi bấm chạy.
        self._blocked_by_user = False

    def running(self):
        return bool(self._thread and self._thread.is_alive())

    @property
    def blocked_by_user(self):
        return self._blocked_by_user

    def state(self):
        if self.running():
            return "running"
        if self._paused_reason:
            return "paused"
        return "idle"

    def start(self, backfill=False, force=False, retry_failed=False, filters=None, skip_failed=False):
        with self._lock:
            if self.running():
                return False
            self._stop.clear()
            self._paused_reason = ""
            self._blocked_by_user = False
            self.status = {"processed": 0, "session_done": 0, "session_failed": 0,
                           "running": True, "phase": "starting", "state": "running",
                           "active_filters": dict(filters or {})}
            self._thread = threading.Thread(
                target=self._run,
                args=(backfill, force, retry_failed, dict(filters or {}), skip_failed),
                daemon=True)
            self._thread.start()
            return True

    def stop(self):
        """Người dùng bấm Dừng: dừng lại và KHÔNG tự chạy lại."""
        self._blocked_by_user = True
        self._paused_reason = ""
        self._stop.set()

    def pause(self, reason="nhường tài nguyên"):
        """Điều phối tạm dừng để nhường máy. Hàng đợi trong CSDL vẫn còn nên
        chỉ cần `start()` lại là chạy tiếp phần dở."""
        if not self.running():
            return
        self._paused_reason = reason or "nhường tài nguyên"
        self._stop.set()

    def resume_allowed(self):
        """Điều phối được phép tự chạy lại không?"""
        return not self._blocked_by_user

    def events_since(self, sequence=0):
        with self._lock:
            return [dict(item) for item in self._events if item["seq"] > int(sequence or 0)]

    def _event(self, level, message, **payload):
        with self._lock:
            self._event_seq += 1
            event = {"seq": self._event_seq, "time": datetime.now().strftime("%H:%M:%S"),
                     "level": level, "message": message, **payload}
            self._events.append(event)
        return event

    def _run(self, backfill, force, retry_failed=False, filters=None, skip_failed=False):
        filters = filters or {}
        db = None
        try:
            db = self.db_factory()
            db.recover_document_queue()
            self.status.update(db.document_stats())
            if retry_failed:
                retried = db.requeue_failed_documents(filters)
                self.status["pending"] = int(self.status.get("pending", 0)) + retried
                message = f"Đã xếp lại {retried:,} CV parsing lỗi để thử lại."
                self._event("info", message)
            if backfill:
                self.status["phase"] = "backfill"
                total, cursor = 0, 0
                while not self._stop.is_set():
                    page = db.enqueue_unparsed_batch(
                        PARSER_VERSION, force, cursor, self.config.parsing_batch_size, filters,
                        skip_failed=skip_failed)
                    total += page['count']
                    cursor = page['last_rowid']
                    if not page['has_more']:
                        break
                    time.sleep(0.05)
                self._event("info", f"Đã xếp hàng {total:,} CV cũ cần bóc tách nội dung.")
                self.status["queued_backfill"] = total
                self.status["pending"] = int(self.status.get("pending", 0)) + total
            done = failed = 0
            self.status["phase"] = "parsing"
            while not self._stop.is_set():
                row = db.claim_document(filters)
                if not row:
                    break
                self.status.update({"current_cv_id": row.get("cv_id", ""),
                                    "current_name": row.get("fullname") or "Ứng viên chưa rõ tên",
                                    "current_filename": row.get("filename", "")})
                path = os.path.abspath(os.path.join(self.config.cv_folder, row['filename']))
                root = os.path.abspath(self.config.cv_folder)
                if os.path.commonpath([root, path]) != root or not os.path.isfile(path):
                    result = _result(path, Path(path).suffix, "error",
                                     error="Không tìm thấy file CV hợp lệ trên ổ đĩa")
                else:
                    result = extract_cv(path, self.config.parsing_ocr_enabled,
                                        self.config.parsing_min_native_chars,
                                        self.config.parsing_max_pages,
                                        self.config.parsing_max_file_mb,
                                        self.config.parsing_max_sheet_cells)
                db.finish_document(row, result)
                name = row.get('fullname') or 'Ứng viên chưa rõ tên'
                quality = round(float(result.get('quality') or 0) * 100)
                if result['status'] == 'done':
                    message = (f"✓ {name} · mã {row['cv_id']} · {row['filename']} → "
                               f"{result.get('method') or 'text'} · {len(result.get('text') or ''):,} ký tự · {quality}%")
                    level = "success"
                else:
                    message = (f"! {name} · mã {row['cv_id']} · {row['filename']} → "
                               f"{result['status']}: {result.get('error') or 'không bóc được nội dung'}")
                    level = "error"
                self._event(level, message, source=row.get('source', ''), account=row.get('account', ''),
                            cv_id=row.get('cv_id', ''), filename=row.get('filename', ''),
                            status=result['status'], method=result.get('method', ''),
                            text_length=len(result.get('text') or ''), quality=quality)
                done += result['status'] == 'done'
                failed += result['status'] != 'done'
                self.status.update({"processed": done + failed, "session_done": done,
                                    "session_failed": failed})
                self.status["pending"] = max(0, int(self.status.get("pending", 0)) - 1)
                if result['status'] == 'done':
                    self.status["done"] = int(self.status.get("done", 0)) + 1
                time.sleep(max(0, int(self.config.parsing_delay_ms)) / 1000.0)
            message = f"Hoàn tất parsing: {done:,} thành công · {failed:,} cần kiểm tra."
            self._event("info", message)
        except Exception as exc:
            self._event("error", f"LỖI PARSING ({type(exc).__name__}): {exc}")
        finally:
            if self._paused_reason:
                phase, state = "paused", "paused"
            elif self._stop.is_set():
                phase, state = "stopped", "idle"
            else:
                phase, state = "complete", "idle"
            self.status.update({"running": False, "phase": phase, "state": state,
                                "paused_reason": self._paused_reason,
                                "blocked_by_user": self._blocked_by_user,
                                "current_cv_id": "", "current_name": "", "current_filename": ""})
            if db:
                db.close()
            # Đẩy trạng thái sang giao diện NGAY khi kết thúc: bản cũ chỉ cập nhật
            # dict trong RAM nên UI chỉ biết ở lần poll kế - đó là lý do nút
            # "Dừng parsing" kẹt lại sau khi đã xong.
            try:
                self._on_finish()
            except Exception:  # noqa: BLE001
                pass
